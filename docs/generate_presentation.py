"""Builds PRESENTATION.pptx from PRESENTATION_CONTENT.md's approved 18-slide outline,
using python-pptx (no Node/LibreOffice available on this machine). Embeds the real
result figures from visualizations/ and the real app screenshots from docs/screenshots/
-- no invented numbers, no stock imagery.

Run: ./venv/Scripts/python.exe docs/generate_presentation.py
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

ROOT = Path(__file__).resolve().parents[1]
VIZ = ROOT / "visualizations"
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "PRESENTATION.pptx"

# ---------------------------------------------------------------- palette --
NAVY = RGBColor(0x1E, 0x27, 0x61)
NAVY_DARK = RGBColor(0x14, 0x1B, 0x46)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
BG_LIGHT = RGBColor(0xF4, 0xF7, 0xFC)
CARD_LIGHT = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x16, 0x21, 0x3E)
MUTED = RGBColor(0x5B, 0x6B, 0x7A)

TITLE_FONT = "Cambria"
BODY_FONT = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, color, line=False, shadow=False, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, text, left, top, width, height, size=16, color=INK, bold=False, italic=False,
             font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing != 1.0:
        p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    return tb


def add_bullets(slide, items, left, top, width, height, size=15, color=INK, font=BODY_FONT,
                 space_after=10, bullet_color=ORANGE, bold_first=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = 1.08
        r = p.add_run()
        r.text = f"•  {item}"
        r.font.size = Pt(size)
        r.font.name = font
        r.font.color.rgb = color
    return tb


def add_page_title(slide, text, kicker=None):
    """Consistent header for every light-background content slide."""
    if kicker:
        add_text(slide, kicker.upper(), Inches(0.7), Inches(0.35), Inches(8), Inches(0.35),
                  size=12, color=ORANGE, bold=True, font=BODY_FONT)
    add_text(slide, text, Inches(0.7), Inches(0.62) if kicker else Inches(0.5), Inches(11.8), Inches(0.9),
              size=32, color=NAVY, bold=True, font=TITLE_FONT)


def add_footer(slide, idx, label):
    add_text(slide, label, Inches(0.7), Inches(7.08), Inches(6), Inches(0.3),
              size=9, color=MUTED, font=BODY_FONT)
    add_text(slide, str(idx), Inches(12.6), Inches(7.08), Inches(0.5), Inches(0.3),
              size=9, color=MUTED, font=BODY_FONT, align=PP_ALIGN.RIGHT)


def add_motif(slide, cx, cy, r=Inches(0.34), colors=(ICE, ORANGE)):
    """Two overlapping circles -- the deck's visual motif (text + image -> fusion)."""
    off = Emu(int(r * 0.55))
    c1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r - off, cy - r, r * 2, r * 2)
    c1.fill.solid(); c1.fill.fore_color.rgb = colors[0]; c1.fill.transparency = 0
    c1.line.fill.background(); c1.shadow.inherit = False
    c2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r + off, cy - r, r * 2, r * 2)
    c2.fill.solid(); c2.fill.fore_color.rgb = colors[1]
    c2.line.fill.background(); c2.shadow.inherit = False
    try:
        c1.fill.fore_color.brightness = 0
        c2.fill.transparency = 0.25
    except Exception:
        pass


def contain(path, max_w, max_h):
    """Returns (w, h) in EMU preserving aspect ratio to fit within max box."""
    im = Image.open(path)
    ratio = im.width / im.height
    box_ratio = max_w / max_h
    if ratio > box_ratio:
        w = max_w
        h = Emu(int(max_w / ratio))
    else:
        h = max_h
        w = Emu(int(max_h * ratio))
    return w, h


def add_picture_contain(slide, path, left, top, max_w, max_h, center=True, frame=True):
    w, h = contain(path, max_w, max_h)
    l = left + (max_w - w) // 2 if center else left
    t = top + (max_h - h) // 2 if center else top
    if frame:
        pad = Pt(4)
        fr = add_rect(slide, l - pad, t - pad, w + 2 * pad, h + 2 * pad, WHITE)
        fr.line.color.rgb = RGBColor(0xDD, 0xE3, 0xEC)
        fr.line.width = Pt(1)
    pic = slide.shapes.add_picture(str(path), l, t, w, h)
    return pic


def stat_tile(slide, left, top, w, h, number, label, number_color=ORANGE):
    card = add_rect(slide, left, top, w, h, CARD_LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    card.line.color.rgb = RGBColor(0xE3, 0xE8, 0xF0)
    card.line.width = Pt(1)
    try:
        card.adjustments[0] = 0.08
    except Exception:
        pass
    add_text(slide, number, left + Inches(0.15), top + Inches(0.18), w - Inches(0.3), Inches(0.75),
              size=30, color=number_color, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
    add_text(slide, label, left + Inches(0.15), top + h - Inches(0.62), w - Inches(0.3), Inches(0.5),
              size=11.5, color=MUTED, font=BODY_FONT, align=PP_ALIGN.CENTER)


def icon_row(slide, left, top, w, h, number, title, body, accent=ORANGE):
    circ_d = Inches(0.55)
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, circ_d, circ_d)
    circ.fill.solid(); circ.fill.fore_color.rgb = accent
    circ.line.fill.background(); circ.shadow.inherit = False
    tf = circ.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(number)
    r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = TITLE_FONT
    add_text(slide, title, left + Inches(0.78), top - Inches(0.02), w - Inches(0.8), Inches(0.4),
              size=16, color=NAVY, bold=True, font=BODY_FONT)
    add_text(slide, body, left + Inches(0.78), top + Inches(0.36), w - Inches(0.8), h - Inches(0.3),
              size=12.5, color=MUTED, font=BODY_FONT, line_spacing=1.1)


def pipeline_box(slide, left, top, w, h, text, fill=NAVY, text_color=WHITE, size=12.5):
    shp = add_rect(slide, left, top, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    try:
        shp.adjustments[0] = 0.15
    except Exception:
        pass
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = text_color; r.font.name = BODY_FONT
    return shp


def pipeline_arrow(slide, left, top, w, h, color=ORANGE):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background(); shp.shadow.inherit = False
    return shp


def card(slide, left, top, w, h, title, items, title_color=NAVY, fill=CARD_LIGHT, border=None):
    c = add_rect(slide, left, top, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    try:
        c.adjustments[0] = 0.04
    except Exception:
        pass
    if border:
        c.line.color.rgb = border
        c.line.width = Pt(1.25)
    add_text(slide, title, left + Inches(0.25), top + Inches(0.2), w - Inches(0.5), Inches(0.5),
              size=17, color=title_color, bold=True, font=TITLE_FONT)
    add_bullets(slide, items, left + Inches(0.25), top + Inches(0.75), w - Inches(0.5), h - Inches(1.0),
                size=12.5, space_after=6)


# ======================================================================
# Slide 1 -- Title
# ======================================================================
s = add_slide()
set_bg(s, NAVY)
add_motif(s, Inches(6.67), Inches(1.55), r=Inches(0.42))
add_text(s, "GROUP 15 · FINAL YEAR PROJECT (CSE) · AKGEC, GHAZIABAD (AKTU)", Inches(1), Inches(2.25),
          Inches(11.3), Inches(0.4), size=13, color=ICE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, "AI-Driven Multimodal Sentiment Analysis for\nMental Health Monitoring on Social Media",
          Inches(1), Inches(2.7), Inches(11.3), Inches(1.9), size=34, color=WHITE, bold=True,
          font=TITLE_FONT, align=PP_ALIGN.CENTER, line_spacing=1.08)
add_text(s, "A Resource-Constrained Comparative Study", Inches(1), Inches(4.35), Inches(11.3), Inches(0.5),
          size=18, color=ORANGE, italic=True, align=PP_ALIGN.CENTER, font=TITLE_FONT)
add_text(s, "Kartikey Tiwari   ·   Akshat Sharma   ·   Karan Sharma   ·   Lalit Sharma",
          Inches(1), Inches(5.3), Inches(11.3), Inches(0.4), size=15, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s, "Supervisor: Dr. Avdhesh Gupta  |  Dept. of CSE  |  Session 2023–2027",
          Inches(1), Inches(5.75), Inches(11.3), Inches(0.4), size=12.5, color=ICE, align=PP_ALIGN.CENTER)

# ======================================================================
# Slide 2 -- Problem
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "The Problem", "Motivation")
rows = [
    ("Half the signal, missed", "A social media post pairs text with an image. Text-only or image-only sentiment analysis each discard the other modality's signal."),
    ("A testable question", "Does combining both actually improve sentiment classification? This project tests it with real, same-dataset experiments rather than assuming the answer."),
    ("Why it matters here", "Framed explicitly as mental-health-related sentiment monitoring support — never diagnosis (Section VIII)."),
]
for i, (t, b) in enumerate(rows):
    icon_row(s, Inches(0.8), Inches(1.9 + i * 1.5), Inches(7.6), Inches(1.3), i + 1, t, b)
add_rect(s, Inches(8.9), Inches(1.9), Inches(3.6), Inches(4.4), WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE).line.color.rgb = ICE
add_motif(s, Inches(10.7), Inches(3.4), r=Inches(0.65))
add_text(s, "Text", Inches(9.3), Inches(4.35), Inches(1.2), Inches(0.35), size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
add_text(s, "Image", Inches(10.9), Inches(4.35), Inches(1.2), Inches(0.35), size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
add_text(s, "Two signals,\none post", Inches(9.1), Inches(4.9), Inches(3.2), Inches(0.8), size=14, color=MUTED,
          italic=True, align=PP_ALIGN.CENTER)
add_footer(s, 2, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 3 -- Motivation
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Motivation", "Why This Research Area")
rows = [
    ("Digital-first awareness", "Mental-health awareness and expression is increasingly happening on social media first."),
    ("A large, real data source", "Social platforms generate a vast, real (not simulated) record of expressed sentiment."),
    ("Support, not replacement", "AI can support mental-health-related monitoring and research — never replace a qualified professional. This non-clinical framing is maintained everywhere: code, UI, paper, and this deck."),
]
for i, (t, b) in enumerate(rows):
    icon_row(s, Inches(0.8), Inches(2.0 + i * 1.55), Inches(11.5), Inches(1.35), i + 1, t, b)
add_footer(s, 3, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 4 -- Existing Approaches (3-column cards)
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Existing Approaches", "Related Work")
cw, gap = Inches(3.75), Inches(0.35)
x0 = Inches(0.7)
card(s, x0, Inches(1.95), cw, Inches(4.5), "Text-based",
     ["NLP / transformers (e.g. MentalBERT)", "Strong language understanding", "Limitation: ignores visual content"],
     border=ICE)
card(s, x0 + cw + gap, Inches(1.95), cw, Inches(4.5), "Image-based",
     ["CNN-based emotion recognition", "Strong visual feature extraction", "Limitation: ignores textual meaning"],
     border=ICE)
card(s, x0 + 2 * (cw + gap), Inches(1.95), cw, Inches(4.5), "Multimodal",
     ["Active, well-populated research area", "Mature fusion-strategy taxonomy exists", "Gap: lacks controlled same-metrics comparison on real social data"],
     border=ORANGE)
add_footer(s, 4, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 5 -- Research Gap
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Research Gap", "What's Missing")
stat_tile(s, Inches(0.8), Inches(2.0), Inches(3.4), Inches(1.9), "100", "pairs — closest prior\nsocial-media comparison")
stat_tile(s, Inches(4.45), Inches(2.0), Inches(3.4), Inches(1.9), "4,511", "real pairs — this\nproject's comparison")
stat_tile(s, Inches(8.1), Inches(2.0), Inches(4.4), Inches(1.9), "CPU-only", "most literature is GPU-scale;\nwe report what this costs")
add_bullets(s, [
    "No located prior work does a controlled, same-dataset, same-metrics text-vs-image-vs-multimodal comparison on adequately-sized real social-media data.",
    "This project addresses that gap directly, and documents the real resource constraints of doing so on consumer CPU-only hardware.",
], Inches(0.8), Inches(4.3), Inches(11.7), Inches(1.8), size=15)
add_footer(s, 5, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 6 -- Objectives
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Objectives", "Research Design")
objs = [
    "Build and compare Text-only, Image-only, and Multimodal models on the same data and metrics",
    "Determine whether — and under what conditions — fusion helps",
    "Analyze modality conflict and errors honestly, including where fusion underperforms",
    "Document resource-constrained reproduction on CPU-only hardware",
]
for i, t in enumerate(objs):
    icon_row(s, Inches(0.8), Inches(2.0 + i * 1.18), Inches(11.5), Inches(1.0), i + 1, t, "")
add_footer(s, 6, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 7 -- Dataset
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Dataset — MVSA-Single", "Research Progress")
add_bullets(s, [
    "4,511 real Twitter image+text pairs (after standard cleaning), 3-class sentiment",
    "Published train/dev/test split (3,611 / 450 / 450) from CLMLF (NAACL 2022) for comparability",
    "Real acquisition challenge: the official distribution link was dead (404) — we found and content-verified a HuggingFace mirror instead",
    "Class imbalance motivates macro-F1 as the primary metric",
], Inches(0.7), Inches(2.0), Inches(5.9), Inches(4.5), size=14.5, space_after=16)
add_picture_contain(s, VIZ / "class_distribution.png", Inches(6.9), Inches(1.9), Inches(5.7), Inches(4.7))
add_footer(s, 7, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 8 -- Proposed Architecture (diagram)
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Proposed Architecture", "Research Design")
bw, bh, y = Inches(3.1), Inches(0.85), Inches(2.2)
pipeline_box(s, Inches(0.8), y, bw, bh, "Text Pipeline\n(cleaning -> embedding)")
pipeline_box(s, Inches(9.4), y, bw, bh, "Image Pipeline\n(resize -> embedding)", fill=NAVY)
pipeline_arrow(s, Inches(4.0), y + Inches(0.22), Inches(0.9), Inches(0.4))
pipeline_arrow(s, Inches(8.4), y + Inches(0.22), Inches(0.9), Inches(0.4))
pipeline_box(s, Inches(5.0), y + Inches(1.35), Inches(3.3), bh, "Fusion\n(feature concatenation)", fill=ORANGE)
# straight connector lines down into the fusion box
conn1 = s.shapes.add_connector(1, Inches(2.35), y + bh, Inches(6.0), y + Inches(1.35))
conn1.line.color.rgb = MUTED; conn1.line.width = Pt(2); conn1.shadow.inherit = False
conn2 = s.shapes.add_connector(1, Inches(10.95), y + bh, Inches(7.0), y + Inches(1.35))
conn2.line.color.rgb = MUTED; conn2.line.width = Pt(2); conn2.shadow.inherit = False
pipeline_arrow(s, Inches(6.35), y + Inches(2.42), Inches(0.65), Inches(0.4))
pipeline_box(s, Inches(5.0), y + Inches(2.75), Inches(3.3), bh, "Classification Head\n(MLP, 3-class softmax)", fill=NAVY)
pipeline_arrow(s, Inches(6.35), y + Inches(3.82), Inches(0.65), Inches(0.4))
pipeline_box(s, Inches(4.6), y + Inches(4.15), Inches(4.1), bh, "Output: Sentiment\n(positive / neutral / negative)", fill=NAVY_DARK)
add_text(s, "Full rationale: PROJECT_ARCHITECTURE.md", Inches(0.8), Inches(6.85), Inches(6), Inches(0.3),
          size=10.5, color=MUTED, italic=True)
add_footer(s, 8, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 9 -- Text Pipeline
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Text Pipeline", "Research Design")
stages = ["Raw tweet\ntext", "Clean\n(URL / mention / emoji)", "Embed\n(TF-IDF or DistilBERT)", "Classify\n(Logistic Regression)"]
bw = Inches(2.55); bh = Inches(1.0); gap = Inches(0.35); y = Inches(2.1)
x = Inches(0.8)
for i, st in enumerate(stages):
    pipeline_box(s, x, y, bw, bh, st, fill=NAVY if i % 2 == 0 else ORANGE)
    x2 = x + bw
    if i < len(stages) - 1:
        pipeline_arrow(s, x2, y + Inches(0.3), gap, Inches(0.4))
    x = x2 + gap
card(s, Inches(0.8), Inches(3.7), Inches(5.7), Inches(2.6), "A1 — TF-IDF baseline",
     ["Unigrams + bigrams, 10k features", "Class-weighted Logistic Regression", "Test: 60.9% accuracy / 52.0% macro-F1"], border=ICE)
card(s, Inches(6.8), Inches(3.7), Inches(5.7), Inches(2.6), "A2 — Frozen DistilBERT",
     ["distilbert-base-uncased, mean-pooled", "Chosen for CPU-only feasibility", "Test: 61.6% accuracy / 54.6% macro-F1"], border=ORANGE)
add_footer(s, 9, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 10 -- Image Pipeline
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Image Pipeline", "Research Design")
stages = ["Raw\nimage", "Resize /\nnormalize", "Embed\n(CNN or ResNet18)", "Classify\n(Logistic Regression)"]
x = Inches(0.8); y = Inches(2.1)
for i, st in enumerate(stages):
    pipeline_box(s, x, y, bw, bh, st, fill=NAVY if i % 2 == 0 else ORANGE)
    x2 = x + bw
    if i < len(stages) - 1:
        pipeline_arrow(s, x2, y + Inches(0.3), gap, Inches(0.4))
    x = x2 + gap
card(s, Inches(0.8), Inches(3.7), Inches(5.7), Inches(2.6), "B1 — CNN from scratch",
     ["Small 4-layer CNN, trained from scratch", "Reduced to 96x96 / 3 epochs (CPU budget)", "Test: 47.6% accuracy / 42.9% macro-F1"], border=ICE)
card(s, Inches(6.8), Inches(3.7), Inches(5.7), Inches(2.6), "B2 — Frozen ResNet18",
     ["ImageNet-pretrained, frozen feature extractor", "Transfer learning over from-scratch training", "Test: 52.7% accuracy / 42.8% macro-F1"], border=ORANGE)
add_footer(s, 10, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 11 -- Multimodal Fusion
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Multimodal Fusion", "Research Design")
pipeline_box(s, Inches(0.8), Inches(2.1), Inches(3.3), Inches(0.9), "Text embedding\n(768-dim)")
pipeline_box(s, Inches(0.8), Inches(3.4), Inches(3.3), Inches(0.9), "Image embedding\n(512-dim)", fill=ORANGE)
pipeline_arrow(s, Inches(4.25), Inches(2.5), Inches(0.7), Inches(0.4))
pipeline_arrow(s, Inches(4.25), Inches(3.8), Inches(0.7), Inches(0.4))
pipeline_box(s, Inches(5.1), Inches(2.75), Inches(3.0), Inches(0.9), "Concatenate\n(1280-dim)", fill=NAVY_DARK)
pipeline_arrow(s, Inches(8.2), Inches(2.95), Inches(0.7), Inches(0.4))
pipeline_box(s, Inches(9.0), Inches(2.75), Inches(3.1), Inches(0.9), "MLP classifier\n(dropout, softmax)", fill=NAVY)
card(s, Inches(0.8), Inches(4.7), Inches(11.3), Inches(1.9), "Primary vs. ablation",
     ["Primary: feature-level (early) concatenation fusion — cheap to train on CPU, standard first fusion strategy in the literature",
      "Ablation: late (decision-level) fusion — independent text/image classifiers' probabilities combined by averaging and by a validation-tuned weight (w_text = 0.80)"],
     border=ICE)
add_footer(s, 11, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 12 -- Experimental Setup
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Experimental Setup", "Research Progress")
tiles = [("Seed = 42", "fixed throughout"), ("CPU-only", "no GPU, 8GB RAM"),
         ("Macro-F1", "prioritized metric"), ("3,611 / 450 / 450", "train / val / test")]
tw = Inches(2.75); tg = Inches(0.3); x = Inches(0.8)
for num, lab in tiles:
    stat_tile(s, x, Inches(2.3), tw, Inches(2.0), num, lab)
    x += tw + tg
add_bullets(s, [
    "Evaluation: accuracy, precision/recall/F1 (macro + weighted), confusion matrix",
    "All numbers read directly from results/*.json, generated by actually-executed scripts under src/",
], Inches(0.8), Inches(4.8), Inches(11.5), Inches(1.5), size=14.5)
add_footer(s, 12, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 13 -- Results
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Results", "Results & Validation")
add_picture_contain(s, VIZ / "model_comparison.png", Inches(0.8), Inches(1.85), Inches(11.7), Inches(4.6))
add_text(s, "Multimodal fusion beats every unimodal model on every metric: 64.4% accuracy / 56.3% macro-F1.",
          Inches(0.8), Inches(6.65), Inches(11.7), Inches(0.5), size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
add_footer(s, 13, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 14 -- Modality Conflict Analysis
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Modality Conflict Analysis", "Results & Validation")
add_picture_contain(s, VIZ / "modality_conflict.png", Inches(0.7), Inches(1.9), Inches(7.3), Inches(4.8))
add_bullets(s, [
    "Text and image models disagree on 53.3% of test posts",
    "On disagreement cases, fusion clearly wins: 57.1% vs. 50.0% (text) vs. 33.3% (image)",
    "On agreement cases, a small regression: fusion 72.9% vs. text 74.8%",
    "An honest finding — not a blanket “multimodal always wins”",
], Inches(8.3), Inches(2.1), Inches(4.3), Inches(4.5), size=13.5, space_after=14)
add_footer(s, 14, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 15 -- Explainability
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Explainability", "Implementation")
rows = [
    ("Text", "Gradient x input token saliency — no extra library required"),
    ("Image", "Grad-CAM via forward/backward hooks on ResNet18's last conv block"),
    ("Multimodal", "Modality-ablation: zero one branch, observe the prediction shift"),
]
for i, (t, b) in enumerate(rows):
    icon_row(s, Inches(0.8), Inches(2.0 + i * 1.3), Inches(11.5), Inches(1.1), i + 1, t, b)
add_text(s, "Exploratory — not a rigorously validated attribution method (documented limitation).",
          Inches(0.8), Inches(6.1), Inches(11.5), Inches(0.5), size=13, color=MUTED, italic=True)
add_footer(s, 15, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 16 -- Application / Dashboard (real screenshots)
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Application & Dashboard", "Implementation — Live Demo")
add_picture_contain(s, SHOTS / "02_predict_result.png", Inches(0.6), Inches(1.75), Inches(5.9), Inches(5.0))
add_picture_contain(s, SHOTS / "03_dashboard.png", Inches(6.85), Inches(1.75), Inches(5.9), Inches(5.0))
add_text(s, "Predict page — real text + image analyzed live", Inches(0.6), Inches(6.85), Inches(5.9), Inches(0.35),
          size=11.5, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
add_text(s, "Research dashboard — model comparison & confusion matrices", Inches(6.85), Inches(6.85), Inches(5.9), Inches(0.35),
          size=11.5, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
add_footer(s, 16, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 17 -- Research Contribution & Limitations
# ======================================================================
s = add_slide(); set_bg(s, BG_LIGHT)
add_page_title(s, "Contribution & Limitations", "Discussion")
card(s, Inches(0.7), Inches(1.95), Inches(5.9), Inches(4.6), "Contribution",
     ["Real, same-conditions unimodal-vs-multimodal comparison on adequately-sized social-media data",
      "Honest modality-conflict analysis, not a blanket “multimodal wins” claim",
      "Documented resource-constrained reproduction (CPU-only, 8GB RAM)"],
     border=ORANGE, title_color=ORANGE)
card(s, Inches(6.75), Inches(1.95), Inches(5.9), Inches(4.6), "Limitations",
     ["Sentiment labels, not mental-health-specific labels (documented proxy)",
      "English-language, 2015–16-era Twitter data only",
      "Small (~1–2%) cross-split duplication in the published split used",
      "From-scratch CNN baseline was compute-constrained"],
     border=ICE)
add_footer(s, 17, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring")

# ======================================================================
# Slide 18 -- Future Work & Conclusion
# ======================================================================
s = add_slide(); set_bg(s, NAVY)
add_motif(s, Inches(6.67), Inches(1.3), r=Inches(0.34))
add_text(s, "CONCLUSION", Inches(1), Inches(1.9), Inches(11.3), Inches(0.4), size=13, color=ORANGE, bold=True,
          align=PP_ALIGN.CENTER)
add_text(s, "Multimodal fusion provides a real, measurable, but conditional\nimprovement — concentrated where modalities disagree.",
          Inches(1), Inches(2.3), Inches(11.3), Inches(1.5), size=24, color=WHITE, bold=True,
          font=TITLE_FONT, align=PP_ALIGN.CENTER, line_spacing=1.15)
add_bullets(s, [
    "Future work: multilingual extension",
    "Future work: cross-modal attention fusion at scale",
    "Future work: validation against a genuinely mental-health-labeled multimodal dataset, if one becomes openly available",
], Inches(2.3), Inches(4.3), Inches(8.7), Inches(2.0), size=14, color=ICE, bullet_color=ORANGE)
add_footer_slide = s
add_text(s, "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring  |  Group 15  |  18",
          Inches(0.7), Inches(7.08), Inches(12), Inches(0.3), size=9, color=ICE, align=PP_ALIGN.CENTER)

prs.save(str(OUT))
print("Wrote", OUT)
