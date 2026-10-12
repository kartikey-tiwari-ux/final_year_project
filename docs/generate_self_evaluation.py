"""Build Viva2_Self_Evaluation_Form.docx: the group's pre-filled self-evaluation against
the six assessment areas named in the Second Project Viva notice (12-13 Oct 2026).

The official template on the CSE Project Central site could not be retrieved, so this
document carries the content; transfer it to the official form if its layout differs.

Run:  ./venv/Scripts/python docs/generate_self_evaluation.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Viva2_Self_Evaluation_Form.docx"

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.6)
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(3)


def para(text, bold=False, size=None, center=False, after=None):
    p = doc.add_paragraph()
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part:
            r = p.add_run(part); r.bold = bold or i % 2 == 1
            if size:
                r.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def table(header, rows, widths, size=10):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = v if isinstance(v, list) else [v]
            for j, line in enumerate(lines):
                p = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                parts = line.split("**")
                for k, part in enumerate(parts):
                    if part:
                        r = p.add_run(part); r.font.size = Pt(size); r.bold = k % 2 == 1
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


para("AJAY KUMAR GARG ENGINEERING COLLEGE, GHAZIABAD", bold=True, size=13, center=True, after=0)
para("Department of Computer Science and Engineering", center=True, after=0)
para("VII Semester Project: SECOND PROJECT VIVA (12–13 October 2026)", bold=True, center=True, after=0)
para("VIVA-II SELF-EVALUATION FORM", bold=True, size=14, center=True, after=8)

table(["Field", "Details"], [
    ["Group No. / Code", "Group 15  (2023-2027/CSE-DS/1/G15)"],
    ["Project Title", "AI-Driven Multimodal Sentiment Analysis for Mental Health Monitoring on Social Media"],
    ["Guide", "Dr. Avdhesh Gupta"],
    ["Members (Roll No.)", ["Kartikey Tiwari (2300271540063)", "Akshat Sharma (2300271540018)",
                           "Karan Sharma (2300271540062)", "Lalit Sharma (2300271540066)"]],
    ["Domain", "AI/ML: NLP, Computer Vision, Multimodal Learning"],
], [4.2, 13.2], size=10.5)

para("A. Self-Assessment Against the Viva-II Criteria", bold=True, size=12, after=4)
table(["#", "Criterion", "What the group has completed (evidence)", "Status", "Self-score"], [
    ["1", "Research Progress",
     ["Problem statement, primary RQ + 4 secondary RQs, 6 objectives defined.",
      "5 of 6 objectives completed; dissemination (O6) in progress.",
      "Evidence: Extended Synopsis Ch. 1, Table 1.1."], "Completed", ""],
    ["2", "Research Design",
     ["Experimental comparative design: same split, seed and metrics for every model.",
      "Architecture: text branch (TF-IDF / DistilBERT), image branch (CNN / ResNet18), concat-MLP and late fusion.",
      "Each choice justified against CPU-only constraints and the literature (Ch. 3, Tables 3.2–3.3)."], "Completed", ""],
    ["3", "Implementation",
     ["Modular Python codebase: data, preprocessing, text, image, fusion, evaluation, explainability.",
      "Streamlit demo + research dashboard (screenshots, live demo).",
      "Datasets: MVSA-Single (4,511), Memotion (3,500). Tools: PyTorch, Transformers, scikit-learn.",
      "32/32 automated tests pass. 7 technical challenges documented with resolutions (Ch. 4)."], "Completed", ""],
    ["4", "Results & Validation",
     ["Study 1 (MVSA): fusion 64.4% acc / 56.3% macro-F1 vs. best text 61.6% / 54.6% and best image 52.7% / 42.8%.",
      "Modality-conflict analysis: the fusion gain is concentrated where text and image disagree (57.1% vs. 50.0%).",
      "Study 2 (memes): blur ablation removes 10.9 pp of fusion accuracy, so the gain comes mainly from embedded-text leakage.",
      "Leakage audit, confusion matrices, explainability (Ch. 5)."], "Completed (multi-seed pending)", ""],
    ["5", "Next Plan",
     ["Multi-seed + significance tests (Oct); OCR-masking ablation and full Memotion run (Nov);",
      "attention fusion as a stretch goal (Nov–Dec); manuscript (Dec); submission (Jan 2027); final report (Feb–Mar 2027). Gantt chart: Ch. 6."],
     "Planned with timeline", ""],
    ["6", "Presentation & Q&A",
     ["Slide deck (PRESENTATION.pptx), IEEE-format paper drafts (2 studies), viva Q&A preparation (VIVA_PREPARATION.md).",
      "Every member can explain their module and the modality-conflict and blur findings."], "Ready", ""],
], [0.7, 2.6, 9.6, 2.4, 2.1], size=9.5)

para("B. Deliverables Brought to the Viva", bold=True, size=12, after=4)
table(["Deliverable", "Status"], [
    ["Extended Synopsis with Research Design and Implementation sections (signed by Guide)", "Prepared: Extended_Synopsis_Viva2.docx / .pdf"],
    ["Viva-II Self-Evaluation Form (signed by Guide)", "This form"],
    ["Presentation slides", "PRESENTATION.pptx"],
    ["Working demonstration", "streamlit run app/streamlit_app.py"],
    ["Research paper drafts (IEEE format)", "research_paper/paper_IEEE.pdf, research_paper_v2/paper_IEEE.pdf"],
], [11, 6.4], size=10)

para("C. Individual Contribution", bold=True, size=12, after=4)
table(["Member", "Contribution so far"], [
    ["Kartikey Tiwari", "Literature review, dataset acquisition (MVSA mirror verification, Memotion), integration"],
    ["Akshat Sharma", "Text preprocessing, TF-IDF and DistilBERT models"],
    ["Karan Sharma", "CNN / ResNet18 image models, fusion models, evaluation and blur ablation"],
    ["Lalit Sharma", "Documentation, research paper, presentation, testing"],
], [4.2, 13.2], size=10)

para("Challenges faced: dead official dataset link (resolved via a verified mirror); CPU-only, 8 GB RAM "
     "(resolved with frozen encoders and cached embeddings); class imbalance (class weighting, macro-F1); "
     "no public multimodal mental-health-labelled dataset (sentiment used as a non-clinical proxy).", after=10)

para("Signatures of Group Members:", bold=True, after=4)
sig = doc.add_table(rows=2, cols=4); sig.style = "Table Grid"
for i, n in enumerate(["Kartikey Tiwari", "Akshat Sharma", "Karan Sharma", "Lalit Sharma"]):
    sig.rows[0].cells[i].text = n; sig.rows[1].cells[i].text = "\n\n"
doc.add_paragraph()
para("Remarks of Guide:", bold=True, after=2)
for _ in range(2):
    para("_" * 80, after=8)
para("Signature of Guide (Dr. Avdhesh Gupta): ______________________          Date: ______________")

doc.save(OUT)
print("wrote", OUT)
